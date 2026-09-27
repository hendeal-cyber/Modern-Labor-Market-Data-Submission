"""Write data/analysis/observations.xlsx: every posting with its application link.

Two sheets: the pay observations (the estimation sample) and every in-scope
posting. Run after any rebuild: PYTHONPATH=src python3 scripts/make_observations_xlsx.py
Requires openpyxl.
"""
import pathlib
import re

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
COLS = {
    "employer": "Employer", "title": "Job title", "Application link": "Application link",
    "industry": "Industry", "role_family": "Role family", "seniority_label": "Seniority",
    "state": "State", "states_listed": "All states listed", "census_region": "Census region",
    "mandate_state": "Pay-transparency mandate (1=yes)", "work_arrangement": "Work arrangement",
    "pay_disclosed": "Pay disclosed (1=yes)", "pay_min": "Pay min (USD/yr)",
    "pay_max": "Pay max (USD/yr)", "pay_midpoint": "Pay midpoint (USD/yr)",
    "hourly_original": "Posted hourly (1=yes)", "yrs_exp_min": "Min years experience",
    "degree_required": "Degree required", "degree_stem": "STEM degree",
    "skill_python_r": "Python/R", "skill_sql": "SQL", "skill_cloud": "Cloud skill",
    "skill_ml_ai": "ML/AI skill", "posted_at": "Posted date",
    "first_seen_run": "First seen", "last_seen_run": "Last seen",
}


def public_link(url: str) -> str:
    """SmartRecruiters rows store the API URL; point at the public job page."""
    m = re.match(r"https://api\.smartrecruiters\.com/v1/companies/([^/]+)/postings/(\d+)", str(url))
    return f"https://jobs.smartrecruiters.com/{m.group(1)}/{m.group(2)}" if m else url


def main() -> None:
    d = pd.read_csv(ROOT / "data" / "analysis" / "postings.csv")
    d["Application link"] = d.url.map(public_link)
    out = d[list(COLS)].rename(columns=COLS)
    out["Posted date"] = out["Posted date"].astype(str).str[:10].replace("nan", "")
    out = out.sort_values(["Employer", "Job title"])
    pay = out[out["Pay disclosed (1=yes)"] == 1]
    path = ROOT / "data" / "analysis" / "observations.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        for name, df in ((f"Pay observations (N={len(pay)})", pay),
                         (f"All in-scope postings ({len(out)})", out)):
            df.to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            ws.freeze_panes = "C2"
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font = Font(bold=True)
            li = list(df.columns).index("Application link") + 1
            for r in range(2, ws.max_row + 1):
                c = ws.cell(r, li)
                if c.value:
                    c.hyperlink, c.value, c.style = c.value, "Apply / view", "Hyperlink"
            for i, col in enumerate(df.columns, 1):
                width = int(df[col].astype(str).str.len().quantile(0.9)) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(45, max(10, len(col) + 2, width))
    print(f"wrote {path} ({len(pay)} pay observations, {len(out)} in scope)")


if __name__ == "__main__":
    main()
